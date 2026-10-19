# @module docx_gen.garantia_docx — DOCX (Word) do Laudo de Avaliação de Garantia.
# Mesmo conteúdo do PDF (services.garantia_conteudo é a fonte única).
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services import garantia_conteudo as C

VERDE = RGBColor(0x0C, 0x33, 0x20)
DOURADO = RGBColor(0xC9, 0xA8, 0x4C)


def _sombra(cell, hexcor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcor)
    tcPr.append(shd)


def _faixa(d: Document, titulo: str):
    t = d.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _sombra(c, "0C3320")
    r = c.paragraphs[0].add_run(titulo)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    d.add_paragraph()


def _kv(d: Document, pares):
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for i, (rot, val) in enumerate(pares):
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Cm(6), Cm(11)
        r = cells[0].paragraphs[0].add_run(str(rot))
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = VERDE
        r2 = cells[1].paragraphs[0].add_run(str(val))
        r2.font.size = Pt(9)
        if i % 2 == 0:
            _sombra(cells[0], "EEF4EF")
            _sombra(cells[1], "EEF4EF")
    d.add_paragraph()


def _img_stream(raw: bytes):
    try:
        from PIL import Image as PILImage
        im = PILImage.open(io.BytesIO(raw)).convert("RGB")
        im.thumbnail((1500, 1500))
        b = io.BytesIO()
        im.save(b, "JPEG", quality=80)
        b.seek(0)
        return b
    except Exception:  # noqa: BLE001
        return None


def generate_garantia_docx(g: dict) -> bytes:
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(C.titulo_laudo(g))
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = VERDE
    for linha in (C.subtitulo(g), f"Laudo nº {g.get('numero') or '—'}"):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(linha).font.size = Pt(11)

    fotos = g.get("_fotos_bytes") or []
    if fotos:
        s = _img_stream(fotos[0][0])
        if s:
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(s, width=Cm(13))
    vp = C.valor_principal(g)
    if vp:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"VALOR DE MERCADO: {C.brl(vp)}")
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = VERDE
    d.add_page_break()

    for n, sec in enumerate(C.secoes(g), 1):
        _faixa(d, f"{n}. {sec['titulo']}")
        for tipo, conteudo in sec["blocos"]:
            if tipo == "kv":
                _kv(d, conteudo)
            else:
                para = d.add_paragraph(conteudo)
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    rt = C.dados_rt(g)
    local = " / ".join([x for x in [g.get("municipio"), g.get("uf")] if x])
    data = C.data_br(g.get("data_avaliacao")) or datetime.now().strftime("%d/%m/%Y")
    for linha, bold in [(f"{local + ', ' if local else ''}{data}.", False), ("", False), ("", False),
                        ("_" * 50, False), ((rt.get("nome") or "").upper(), True),
                        ("Avaliador — " + " · ".join(rt.get("registros_linhas") or []) if rt.get("registros_linhas") else "", False)]:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if linha:
            p.add_run(linha).bold = bold

    if fotos:
        d.add_page_break()
        _faixa(d, "ANEXO I — RELATÓRIO FOTOGRÁFICO")
        for raw, leg in fotos:
            s = _img_stream(raw)
            if not s:
                continue
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(s, width=Cm(12))
            c = d.add_paragraph(leg)
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
